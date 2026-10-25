"""Import d'utilisateurs depuis un fichier CSV ou Excel (.xlsx)."""

import csv
import io
import re
import secrets
import unicodedata

from django.contrib.auth.tokens import default_token_generator
from django.core.exceptions import ValidationError
from django.core.mail import send_mail
from django.core.validators import validate_email
from django.db import transaction
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode

from .forms import creer_profil
from .models import Utilisateur

COLONNES = ["prenom", "nom", "email", "role", "nom_utilisateur", "niveau_ou_specialite"]
ROLES = {
    "etudiant": Utilisateur.Role.ETUDIANT, "étudiant": Utilisateur.Role.ETUDIANT,
    "enseignant": Utilisateur.Role.ENSEIGNANT, "professeur": Utilisateur.Role.ENSEIGNANT,
}
TAILLE_MAX = 2 * 1024 * 1024  # 2 Mo
LIGNES_MAX = 1000


class ImportErreur(Exception):
    """Fichier inutilisable (format, colonnes manquantes…)."""


def _normaliser(texte):
    """« Prénom » -> « prenom » : minuscules, sans accents ni espaces."""
    texte = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9_]", "", texte.strip().lower().replace(" ", "_"))


def lire_fichier(fichier):
    """Renvoie la liste des lignes sous forme de dictionnaires {colonne: valeur}."""
    if fichier.size > TAILLE_MAX:
        raise ImportErreur("Fichier trop volumineux (2 Mo maximum).")
    nom = fichier.name.lower()

    if nom.endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            feuille = load_workbook(fichier, read_only=True, data_only=True).active
        except Exception:
            raise ImportErreur("Fichier Excel illisible.")
        lignes = [["" if v is None else str(v).strip() for v in ligne] for ligne in feuille.iter_rows(values_only=True)]
    elif nom.endswith(".csv"):
        contenu = fichier.read()
        for encodage in ("utf-8-sig", "cp1252"):  # Excel enregistre souvent les CSV en cp1252
            try:
                texte = contenu.decode(encodage)
                break
            except UnicodeDecodeError:
                continue
        separateur = ";" if texte.split("\n", 1)[0].count(";") >= texte.split("\n", 1)[0].count(",") else ","
        lignes = [[v.strip() for v in ligne] for ligne in csv.reader(io.StringIO(texte), delimiter=separateur)]
    else:
        raise ImportErreur("Format non pris en charge : utilisez un fichier .csv ou .xlsx.")

    lignes = [l for l in lignes if any(l)]
    if len(lignes) < 2:
        raise ImportErreur("Le fichier ne contient aucune ligne de données.")
    if len(lignes) - 1 > LIGNES_MAX:
        raise ImportErreur(f"Trop de lignes ({LIGNES_MAX} maximum par import).")

    entetes = [_normaliser(e) for e in lignes[0]]
    manquantes = [c for c in ("prenom", "nom", "email") if c not in entetes]
    if manquantes:
        raise ImportErreur("Colonnes obligatoires manquantes : " + ", ".join(manquantes)
                           + ". Téléchargez le modèle pour voir le format attendu.")
    return [dict(zip(entetes, ligne)) for ligne in lignes[1:]]


def _nom_utilisateur_libre(prenom, nom, deja_pris):
    base = (_normaliser(prenom)[:1] + _normaliser(nom))[:20] or "utilisateur"
    candidat, numero = base, 1
    while candidat in deja_pris or Utilisateur.objects.filter(username__iexact=candidat).exists():
        numero += 1
        candidat = f"{base}{numero}"
    return candidat


def importer(lignes, role_par_defaut=Utilisateur.Role.ETUDIANT):
    """Crée les comptes valides. Renvoie (comptes créés, erreurs [(numéro de ligne, message)])."""
    crees, erreurs, emails_vus, noms_vus = [], [], set(), set()

    for numero, ligne in enumerate(lignes, start=2):  # ligne 1 = en-têtes
        prenom, nom = ligne.get("prenom", "").strip(), ligne.get("nom", "").strip()
        email = ligne.get("email", "").strip().lower()
        if not prenom or not nom or not email:
            erreurs.append((numero, "Prénom, nom et e-mail sont obligatoires."))
            continue
        try:
            validate_email(email)
        except ValidationError:
            erreurs.append((numero, f"E-mail invalide : {email}"))
            continue
        if email in emails_vus or Utilisateur.objects.filter(email__iexact=email).exists():
            erreurs.append((numero, f"E-mail déjà utilisé : {email}"))
            continue

        texte_role = _normaliser(ligne.get("role", ""))
        if texte_role and texte_role not in {_normaliser(r) for r in ROLES}:
            erreurs.append((numero, f"Rôle inconnu : « {ligne.get('role')} » (Etudiant ou Enseignant)."))
            continue
        role = next((v for k, v in ROLES.items() if _normaliser(k) == texte_role), role_par_defaut)

        nom_utilisateur = ligne.get("nom_utilisateur", "").strip()
        if nom_utilisateur:
            if not re.fullmatch(r"[\w.@+-]{1,150}", nom_utilisateur):
                erreurs.append((numero, f"Nom d'utilisateur invalide : {nom_utilisateur}"))
                continue
            if nom_utilisateur.lower() in noms_vus or Utilisateur.objects.filter(username__iexact=nom_utilisateur).exists():
                erreurs.append((numero, f"Nom d'utilisateur déjà pris : {nom_utilisateur}"))
                continue
        else:
            nom_utilisateur = _nom_utilisateur_libre(prenom, nom, noms_vus)

        with transaction.atomic():
            user = Utilisateur(username=nom_utilisateur, first_name=prenom, last_name=nom, email=email,
                               role=role, email_verifie=True)
            # Mot de passe aléatoire inconnu de tous : l'utilisateur choisit le sien grâce au lien reçu
            user.set_password(secrets.token_urlsafe(24))
            user.save()
            creer_profil(user)
            complement = ligne.get("niveau_ou_specialite", "").strip()
            if complement and hasattr(user, "etudiant"):
                user.etudiant.niveau = complement
                user.etudiant.save()
            elif complement and hasattr(user, "enseignant"):
                user.enseignant.specialite = complement
                user.enseignant.save()

        emails_vus.add(email)
        noms_vus.add(nom_utilisateur.lower())
        crees.append(user)
    return crees, erreurs


def envoyer_invitation(request, user):
    """E-mail de bienvenue avec un lien pour choisir son mot de passe."""
    lien = request.build_absolute_uri(reverse("mdp-reinitialiser", args=[
        urlsafe_base64_encode(force_bytes(user.pk)), default_token_generator.make_token(user),
    ]))
    message = render_to_string("accounts/emails/invitation.txt", {"user": user, "lien": lien,
                                                                  "connexion": request.build_absolute_uri(reverse("connexion"))})
    send_mail("Bienvenue sur EduVox – Activez votre compte", message, None, [user.email], fail_silently=True)